#!/usr/bin/env python3
"""
CleanWave Complete - Full Factory Reset Assistant
All features: AI batching, resume, parallel hashing, SQLite, undo, etc.
"""

import os
import sys
import json
import hashlib
import shutil
import argparse
import datetime
import time
import fnmatch
import sqlite3
from pathlib import Path
from typing import Dict, List, Tuple, Optional, Any
from dataclasses import dataclass, asdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from collections import defaultdict

# Dependencies with graceful fallback
try:
    import requests
    from plyer import notification
except ImportError as e:
    print(f"Missing required dependency: {e}")
    print("Run: pip install requests plyer")
    sys.exit(1)

try:
    from rich.console import Console
    from rich.progress import Progress, SpinnerColumn, TextColumn, BarColumn, TimeRemainingColumn, ProgressColumn
    from rich.table import Table
    from rich.panel import Panel
    from rich.prompt import Confirm, Prompt
    from rich.live import Live
    from rich.layout import Layout
    RICH_AVAILABLE = True
    console = Console()
except ImportError:
    RICH_AVAILABLE = False
    console = None
    print("Tip: Install 'rich' for better UI: pip install rich")

# Optional content extraction
try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import pdfplumber
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    import yaml
    YAML_AVAILABLE = True
except ImportError:
    YAML_AVAILABLE = False
    print("Warning: 'yaml' not installed. Using JSON config. pip install pyyaml")

# Optional Ollama
try:
    import ollama
    OLLAMA_AVAILABLE = True
except ImportError:
    OLLAMA_AVAILABLE = False

# ======================= CONFIGURATION =======================
GROQ_API_KEY = os.environ.get("GROQ_API_KEY")
GROQ_API_URL = "https://api.groq.com/openai/v1/chat/completions"

DEFAULT_CONFIDENCE = 0.75
DUPLICATE_HASH_BLOCK_SIZE = 65536
MAX_TEXT_PREVIEW_CHARS = 500

# Extension groups
TEMP_EXTENSIONS = {'.tmp', '.temp', '.cache', '.cached', '.log', '.bak', '.old', '.swp', '.~', '.part'}
INSTALLER_EXTENSIONS = {'.exe', '.msi', '.dmg', '.pkg', '.deb', '.rpm', '.appimage', '.run'}
IMPORTANT_EXTENSIONS = {'.docx', '.xlsx', '.pptx', '.pdf', '.py', '.c', '.cpp', '.java', '.js', '.html', '.css', '.md', '.txt'}

DELETE_AFTER_DAYS_DOWNLOADS = 30
DELETE_AFTER_DAYS_CACHE = 14

# ======================= DATA CLASSES =======================
@dataclass
class FileInfo:
    path: Path
    size: int
    mtime: float
    ext: str
    hash_val: Optional[str] = None

@dataclass
class Decision:
    deletable: bool
    confidence: float
    category: str
    reason: str
    suggested_path: Optional[str] = None
    suggested_name: Optional[str] = None

@dataclass
class BatchResult:
    decisions: List[Optional[Decision]]
    failed_indices: List[int]

# ======================= DATABASE =======================
class CleanWaveDB:
    def __init__(self, db_path: Path):
        self.db_path = db_path
        self.conn = None
        self.init_db()
    
    def init_db(self):
        self.conn = sqlite3.connect(str(self.db_path), check_same_thread=False)
        self.conn.execute("PRAGMA journal_mode=WAL")
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS file_hashes (
                path TEXT PRIMARY KEY,
                size INTEGER,
                mtime REAL,
                hash TEXT,
                last_scanned REAL
            )
        """)
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS checkpoint (
                id INTEGER PRIMARY KEY CHECK (id = 1),
                last_path TEXT,
                processed_count INTEGER,
                batch_number INTEGER,
                timestamp REAL
            )
        """)
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS moved_files (
                original_path TEXT PRIMARY KEY,
                moved_to TEXT,
                decision_category TEXT,
                moved_at REAL,
                restored BOOLEAN DEFAULT 0
            )
        """)
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS scan_history (
                scan_id INTEGER PRIMARY KEY AUTOINCREMENT,
                start_time REAL,
                end_time REAL,
                files_scanned INTEGER,
                files_moved INTEGER,
                config_snapshot TEXT
            )
        """)
        self.conn.commit()
    
    def get_cached_hash(self, path: Path, size: int, mtime: float) -> Optional[str]:
        cursor = self.conn.execute(
            "SELECT hash FROM file_hashes WHERE path = ? AND size = ? AND mtime = ?",
            (str(path), size, mtime)
        )
        row = cursor.fetchone()
        return row[0] if row else None
    
    def store_hash(self, path: Path, size: int, mtime: float, hash_val: str):
        self.conn.execute(
            "INSERT OR REPLACE INTO file_hashes (path, size, mtime, hash, last_scanned) VALUES (?, ?, ?, ?, ?)",
            (str(path), size, mtime, hash_val, time.time())
        )
        self.conn.commit()
    
    def was_moved(self, path: Path) -> bool:
        cursor = self.conn.execute("SELECT 1 FROM moved_files WHERE original_path = ? AND restored = 0", (str(path),))
        return cursor.fetchone() is not None
    
    def record_moved(self, path: Path, dest: Path, category: str):
        self.conn.execute(
            "INSERT OR REPLACE INTO moved_files (original_path, moved_to, decision_category, moved_at) VALUES (?, ?, ?, ?)",
            (str(path), str(dest), category, time.time())
        )
        self.conn.commit()
    
    def record_restored(self, path: Path):
        self.conn.execute("UPDATE moved_files SET restored = 1 WHERE original_path = ?", (str(path),))
        self.conn.commit()
    
    def get_checkpoint(self) -> Optional[Tuple[str, int, int]]:
        cursor = self.conn.execute("SELECT last_path, processed_count, batch_number FROM checkpoint WHERE id = 1")
        row = cursor.fetchone()
        return (row[0], row[1], row[2]) if row else None
    
    def save_checkpoint(self, last_path: str, processed_count: int, batch_number: int):
        self.conn.execute(
            "INSERT OR REPLACE INTO checkpoint (id, last_path, processed_count, batch_number, timestamp) VALUES (1, ?, ?, ?, ?)",
            (last_path, processed_count, batch_number, time.time())
        )
        self.conn.commit()
    
    def begin_scan(self, config: dict) -> int:
        cursor = self.conn.execute(
            "INSERT INTO scan_history (start_time, config_snapshot) VALUES (?, ?)",
            (time.time(), json.dumps(config))
        )
        self.conn.commit()
        return cursor.lastrowid
    
    def end_scan(self, scan_id: int, files_scanned: int, files_moved: int):
        self.conn.execute(
            "UPDATE scan_history SET end_time = ?, files_scanned = ?, files_moved = ? WHERE scan_id = ?",
            (time.time(), files_scanned, files_moved, scan_id)
        )
        self.conn.commit()
    
    def close(self):
        if self.conn:
            self.conn.close()

# ======================= WHITELIST =======================
class Whitelist:
    def __init__(self, patterns: List[str], paths: List[str]):
        self.patterns = patterns
        self.paths = {Path(p).expanduser().resolve() for p in paths}
    
    def is_whitelisted(self, path: Path) -> bool:
        abs_path = path.resolve()
        if abs_path in self.paths:
            return True
        for parent in abs_path.parents:
            if parent in self.paths:
                return True
        for pattern in self.patterns:
            if fnmatch.fnmatch(path.name, pattern):
                return True
            if fnmatch.fnmatch(str(path), pattern):
                return True
        return False

# ======================= CONTENT EXTRACTION =======================
def extract_content(file_path: Path) -> str:
    """Extract text from various file types for AI analysis."""
    ext = file_path.suffix.lower()
    
    try:
        # Text files
        if ext in {'.txt', '.md', '.py', '.js', '.html', '.css', '.json', '.xml', '.csv', '.log'}:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read(MAX_TEXT_PREVIEW_CHARS)
        
        # Word documents
        elif ext == '.docx' and DOCX_AVAILABLE:
            doc = Document(file_path)
            text = ' '.join([para.text for para in doc.paragraphs[:15]])
            return text[:MAX_TEXT_PREVIEW_CHARS]
        
        # PDFs
        elif ext == '.pdf' and PDF_AVAILABLE:
            with pdfplumber.open(file_path) as pdf:
                text = ''
                for page in pdf.pages[:3]:
                    text += page.extract_text() or ''
            return text[:MAX_TEXT_PREVIEW_CHARS]
        
        # Excel files
        elif ext in {'.xlsx', '.xls'}:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                text = ''
                for sheet in wb.worksheets[:2]:
                    for row in sheet.iter_rows(values_only=True, max_row=20):
                        text += ' '.join(str(cell) for cell in row if cell)
                return text[:MAX_TEXT_PREVIEW_CHARS]
            except:
                pass
    
    except Exception:
        pass
    
    return ""

# ======================= OS HELPERS =======================
def get_os_exclusions() -> List[str]:
    if sys.platform == "darwin":
        return [
            "/System", "/Library", "/Applications", "/usr", "/bin", "/sbin",
            "/private", "/Volumes", "/Network", "/cores", "/dev"
        ]
    elif sys.platform == "win32":
        return [
            "C:\\Windows", "C:\\Program Files", "C:\\Program Files (x86)",
            "C:\\ProgramData", "C:\\$Recycle.Bin", "C:\\System Volume Information",
            "C:\\Recovery", "C:\\PerfLogs"
        ]
    else:
        return ["/proc", "/sys", "/dev", "/boot"]

def check_mount(path: Path) -> Tuple[bool, str]:
    """Check if path is mounted and writable."""
    if not path.exists():
        return False, "Path does not exist"
    try:
        test_file = path / f".cleanwave_test_{os.getpid()}"
        test_file.write_text("test")
        test_file.unlink()
        return True, "OK"
    except Exception as e:
        return False, str(e)

def check_disk_space(path: Path, required_gb: float = 5.0) -> Tuple[bool, float]:
    """Check free disk space."""
    try:
        free = shutil.disk_usage(path).free
        free_gb = free / (1024**3)
        return free_gb >= required_gb, free_gb
    except:
        return False, 0.0

# ======================= RULES ENGINE =======================
def quick_rule_deletable(file_info: FileInfo, is_downloads: bool = False) -> Optional[Decision]:
    """Fast rule-based classification - no AI needed."""
    ext = file_info.ext.lower()
    age_days = (time.time() - file_info.mtime) / 86400
    
    # Empty files
    if file_info.size == 0:
        return Decision(True, 1.0, "zero_byte", "File is completely empty")
    
    # Temporary files
    if ext in TEMP_EXTENSIONS:
        return Decision(True, 0.95, "temporary", f"Temporary file type: {ext}")
    
    # Old logs
    if ext == '.log' and age_days > 7:
        return Decision(True, 0.90, "old_log", f"Log file older than {age_days:.0f} days")
    
    # Cache files
    if 'cache' in str(file_info.path).lower() and age_days > DELETE_AFTER_DAYS_CACHE:
        return Decision(True, 0.85, "cache", f"Cache file older than {DELETE_AFTER_DAYS_CACHE} days")
    
    # Old downloads
    if is_downloads and age_days > DELETE_AFTER_DAYS_DOWNLOADS:
        return Decision(True, 0.80, "old_download", f"Download older than {DELETE_AFTER_DAYS_DOWNLOADS} days")
    
    # Unused installers
    if ext in INSTALLER_EXTENSIONS and age_days > 7:
        return Decision(True, 0.75, "old_installer", f"Installer not used for {age_days:.0f} days")
    
    # System temp patterns
    basename = file_info.path.name.lower()
    if basename.startswith("~$") or basename.endswith(".tmp") or basename == "thumbs.db":
        return Decision(True, 0.98, "system_temp", "Auto-generated temporary file")
    
    return None

def find_duplicates_parallel(files: List[FileInfo], db: CleanWaveDB, max_workers: int = 4) -> Dict[str, List[FileInfo]]:
    """Find duplicate files using parallel hash computation."""
    
    def compute_hash_with_cache(f: FileInfo) -> Tuple[FileInfo, Optional[str]]:
        cached = db.get_cached_hash(f.path, f.size, f.mtime)
        if cached:
            return f, cached
        hash_val = compute_file_hash(f.path)
        if hash_val:
            db.store_hash(f.path, f.size, f.mtime, hash_val)
        return f, hash_val
    
    # Compute hashes in parallel
    hash_map = defaultdict(list)
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = [executor.submit(compute_hash_with_cache, f) for f in files]
        for future in as_completed(futures):
            file_info, hash_val = future.result()
            if hash_val:
                file_info.hash_val = hash_val
                hash_map[hash_val].append(file_info)
    
    # Return only duplicates (length > 1)
    return {h: files for h, files in hash_map.items() if len(files) > 1}

def compute_file_hash(path: Path) -> Optional[str]:
    """Compute SHA256 hash of file."""
    hasher = hashlib.sha256()
    try:
        with open(path, 'rb') as f:
            while chunk := f.read(DUPLICATE_HASH_BLOCK_SIZE):
                hasher.update(chunk)
        return hasher.hexdigest()
    except (OSError, PermissionError):
        return None

# ======================= AI ENGINE =======================
def build_ai_prompt_batch(file_infos: List[FileInfo], contents: List[str]) -> str:
    """Build batch prompt for multiple files."""
    prompt = """You are a file cleanup assistant. For each file below, decide if it can be safely deleted.
Output a JSON array with one object per file in the same order.
Each object: {"deletable": bool, "confidence": 0.0-1.0, "category": string, "reason": string}

Files:
"""
    for i, (f, content) in enumerate(zip(file_infos, contents)):
        age_days = (time.time() - f.mtime) / 86400
        prompt += f"""
--- FILE {i} ---
Name: {f.path.name}
Path: {f.path}
Extension: {f.ext or 'none'}
Size KB: {f.size / 1024:.2f}
Age days: {age_days:.1f}
Preview: {content[:300]}
"""
    return prompt

def call_groq_batch(file_infos: List[FileInfo], contents: List[str], retry_count: int = 2) -> BatchResult:
    """Send batch to Groq, retry failed ones."""
    if not GROQ_API_KEY:
        return BatchResult([None] * len(file_infos), list(range(len(file_infos))))
    
    prompt = build_ai_prompt_batch(file_infos, contents)
    headers = {"Authorization": f"Bearer {GROQ_API_KEY}", "Content-Type": "application/json"}
    payload = {
        "model": "llama3-8b-8192",
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.2,
        "response_format": {"type": "json_object"}
    }
    
    for attempt in range(retry_count + 1):
        try:
            resp = requests.post(GROQ_API_URL, json=payload, headers=headers, timeout=60)
            resp.raise_for_status()
            data = resp.json()
            result = json.loads(data["choices"][0]["message"]["content"])
            
            # Parse response - handle both array and object formats
            if isinstance(result, list):
                decisions = []
                failed = []
                for i, item in enumerate(result):
                    if item and 'deletable' in item:
                        decisions.append(Decision(
                            deletable=item['deletable'],
                            confidence=item.get('confidence', 0.5),
                            category=item.get('category', 'unknown'),
                            reason=item.get('reason', 'AI decision')
                        ))
                    else:
                        decisions.append(None)
                        failed.append(i)
                return BatchResult(decisions, failed)
            else:
                # Single file response format
                return BatchResult([Decision(
                    deletable=result.get('deletable', False),
                    confidence=result.get('confidence', 0.5),
                    category=result.get('category', 'unknown'),
                    reason=result.get('reason', 'AI decision')
                )], [])
        
        except Exception as e:
            if attempt == retry_count:
                return BatchResult([None] * len(file_infos), list(range(len(file_infos))))
            time.sleep(2 ** attempt)  # Exponential backoff
    
    return BatchResult([None] * len(file_infos), list(range(len(file_infos))))

def call_local_ollama(file_info: FileInfo, content: str) -> Optional[Decision]:
    """Fallback to local Ollama if available."""
    if not OLLAMA_AVAILABLE:
        return None
    
    try:
        age_days = (time.time() - file_info.mtime) / 86400
        prompt = f"""Decide if this file can be safely deleted. Output ONLY JSON.
File: {file_info.path.name}
Extension: {file_info.ext}
Size KB: {file_info.size/1024:.2f}
Age days: {age_days:.1f}
Content preview: {content[:300]}

Output: {{"deletable": true/false, "confidence": 0.0-1.0, "category": "string", "reason": "string"}}"""
        
        response = ollama.chat(model='llama3.2:3b', messages=[{'role': 'user', 'content': prompt}])
        result = json.loads(response['message']['content'])
        return Decision(
            deletable=result.get('deletable', False),
            confidence=result.get('confidence', 0.5),
            category=result.get('category', 'unknown'),
            reason=result.get('reason', 'Local AI decision')
        )
    except:
        return None

# ======================= FILE MOVEMENT =======================
def move_to_review_folder(file_info: FileInfo, base_dest: Path, dry_run: bool) -> Optional[Path]:
    """Move file to review folder, preserving relative path structure."""
    src = file_info.path
    home = Path.home()
    
    # Create relative path from home
    try:
        rel = src.relative_to(home)
    except ValueError:
        # File outside home (e.g., other drive)
        if sys.platform == "win32":
            drive_letter = src.drive[0] if src.drive else "C"
            rel = Path(f"drive_{drive_letter}") / src.relative_to(src.anchor)
        else:
            rel = Path("root") / src.relative_to(src.anchor)
    
    dest = base_dest / rel
    dest.parent.mkdir(parents=True, exist_ok=True)
    
    if dry_run:
        return dest
    
    # Handle collisions
    if dest.exists():
        stem = dest.stem
        suffix = dest.suffix
        counter = 1
        while dest.exists():
            dest = dest.with_name(f"{stem}_duplicate_{counter}{suffix}")
            counter += 1
    
    try:
        shutil.move(str(src), str(dest))
        return dest
    except Exception as e:
        console.print(f"[red]Error moving {src}: {e}[/red]")
        return None

def generate_restore_script(moved_files: List[Tuple[Path, Path]], output_dir: Path):
    """Generate script to restore moved files to original locations."""
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    
    if sys.platform == "win32":
        script_path = output_dir / f"restore_files_{timestamp}.bat"
        with open(script_path, 'w') as f:
            f.write("@echo off\n")
            f.write(f"REM CleanWave Restore Script - Generated {datetime.datetime.now()}\n")
            f.write("echo Restoring files...\n\n")
            for orig, dest in moved_files:
                f.write(f'if exist "{dest}" (\n')
                f.write(f'  move "{dest}" "{orig}"\n')
                f.write(f'  echo Restored: {orig}\n')
                f.write(f') else (\n')
                f.write(f'  echo Warning: {dest} not found\n')
                f.write(f')\n')
            f.write('\necho Done!\n')
            f.write('pause\n')
    else:
        script_path = output_dir / f"restore_files_{timestamp}.sh"
        with open(script_path, 'w') as f:
            f.write("#!/bin/bash\n")
            f.write(f"# CleanWave Restore Script - Generated {datetime.datetime.now()}\n")
            f.write("echo 'Restoring files...'\n\n")
            for orig, dest in moved_files:
                f.write(f'if [ -f "{dest}" ]; then\n')
                f.write(f'  mv "{dest}" "{orig}"\n')
                f.write(f'  echo "Restored: {orig}"\n')
                f.write(f'else\n')
                f.write(f'  echo "Warning: {dest} not found"\n')
                f.write(f'fi\n')
            f.write('\necho "Done!"\n')
        script_path.chmod(0o755)
    
    console.print(f"[green]✓ Restore script generated: {script_path}[/green]")
    return script_path

# ======================= COLLECTION & PROCESSING =======================
def collect_files(scan_dirs: List[str], exclusions: List[str], db: CleanWaveDB, 
                  skip_moved: bool, resume_path: Optional[str] = None) -> List[FileInfo]:
    """Walk directories and collect file info, with resume support."""
    files = []
    resume_mode = resume_path is not None
    found_resume_point = not resume_mode
    
    for start_dir in scan_dirs:
        start_path = Path(start_dir).expanduser()
        if not start_path.exists():
            console.print(f"[yellow]Warning: {start_dir} does not exist, skipping[/yellow]")
            continue
        
        for root, dirs, filenames in os.walk(start_path):
            root_path = Path(root)
            
            # Filter excluded directories
            dirs[:] = [d for d in dirs if not any(
                (root_path / d).is_relative_to(Path(excl).expanduser()) 
                for excl in exclusions
            )]
            
            for fname in filenames:
                fpath = root_path / fname
                
                # Skip excluded paths
                if any(fpath.is_relative_to(Path(excl).expanduser()) for excl in exclusions):
                    continue
                
                # Skip already moved files
                if skip_moved and db.was_moved(fpath):
                    continue
                
                # Resume support - skip until we hit the resume point
                if resume_mode and not found_resume_point:
                    if str(fpath) == resume_path:
                        found_resume_point = True
                    continue
                
                try:
                    stat = fpath.stat()
                    files.append(FileInfo(
                        path=fpath,
                        size=stat.st_size,
                        mtime=stat.st_mtime,
                        ext=fpath.suffix.lower()
                    ))
                except (OSError, PermissionError):
                    continue
    
    return files

def process_files_advanced(files: List[FileInfo], 
                           deletion_base: Path,
                           low_conf_base: Path,
                           db: CleanWaveDB,
                           whitelist: Whitelist,
                           config: dict,
                           dry_run: bool,
                           safe_mode: bool,
                           resume: bool) -> Tuple[List, List, List]:
    """Main processing loop with AI batching, resume, and duplicate detection."""
    
    moved_deletion = []
    moved_low = []
    restore_list = []
    
    # Find resume point
    start_idx = 0
    batch_number = 0
    if resume:
        checkpoint = db.get_checkpoint()
        if checkpoint:
            last_path, processed_count, batch_number = checkpoint
            for i, f in enumerate(files):
                if str(f.path) == last_path:
                    start_idx = i + 1
                    break
            console.print(f"[yellow]Resuming from file {start_idx}/{len(files)}[/yellow]")
    
    # Initialize progress display
    if RICH_AVAILABLE:
        from rich.progress import Progress
        progress = Progress(
            SpinnerColumn(),
            TextColumn("[progress.description]{task.description}"),
            BarColumn(),
            TextColumn("[progress.percentage]{task.percentage:>3.0f}%"),
            TimeRemainingColumn(),
            console=console
        )
        task = progress.add_task("[cyan]Processing files...", total=len(files) - start_idx)
        progress.start()
    else:
        progress = None
    
    # First pass: duplicate detection on all files (parallel)
    console.print("[cyan]🔍 Finding duplicates...[/cyan]")
    duplicate_groups = find_duplicates_parallel(files[start_idx:], db, config.get('parallel_hash_workers', 4))
    
    duplicate_decisions = {}
    for hash_val, dup_files in duplicate_groups.items():
        # Keep the oldest file (by modification time)
        dup_files.sort(key=lambda x: x.mtime)
        keeper = dup_files[0]
        for f in dup_files[1:]:
            duplicate_decisions[f.path] = Decision(
                True, 0.99, "duplicate", 
                f"Duplicate of {keeper.path.name} (same content)"
            )
    
    # Process files in batches for AI
    batch_size = config.get('batch_size', 6)
    ai_queue = []
    
    for idx, file_info in enumerate(files[start_idx:], start=start_idx):
        if progress:
            progress.update(task, advance=1, description=f"[cyan]{file_info.path.name[:60]}")
        
        # Checkpoint every N files
        if (idx + 1) % CHECKPOINT_INTERVAL == 0:
            db.save_checkpoint(str(file_info.path), idx + 1, batch_number)
        
        # Whitelist check
        if whitelist.is_whitelisted(file_info.path):
            continue
        
        # Duplicate check
        if file_info.path in duplicate_decisions:
            decision = duplicate_decisions[file_info.path]
            dest = move_to_review_folder(file_info, deletion_base, dry_run) if not safe_mode else None
            if not safe_mode and dest:
                moved_deletion.append((file_info, decision, dest))
                db.record_moved(file_info.path, dest, decision.category)
                restore_list.append((file_info.path, dest))
            elif safe_mode:
                moved_deletion.append((file_info, decision, None))
            continue
        
        # Rule-based check
        is_downloads = "downloads" in str(file_info.path).lower()
        rule_decision = quick_rule_deletable(file_info, is_downloads)
        
        if rule_decision and rule_decision.confidence >= config.get('confidence_threshold', 0.75):
            dest = move_to_review_folder(file_info, deletion_base, dry_run) if not safe_mode else None
            if not safe_mode and dest:
                moved_deletion.append((file_info, rule_decision, dest))
                db.record_moved(file_info.path, dest, rule_decision.category)
                restore_list.append((file_info.path, dest))
            elif safe_mode:
                moved_deletion.append((file_info, rule_decision, None))
            continue
        
        # Queue for AI processing
        if config.get('ai_enabled', True):
            ai_queue.append(file_info)
            
            # Process batch when full
            if len(ai_queue) >= batch_size:
                batch_result = process_ai_batch(ai_queue, db, config, low_conf_base, 
                                                deletion_base, dry_run, safe_mode)
                for result in batch_result:
                    if result:
                        moved_deletion.append(result) if result[1].deletable else moved_low.append(result)
                        if not dry_run and not safe_mode:
                            restore_list.append((result[0].path, result[2]))
                ai_queue = []
                batch_number += 1
    
    # Process remaining AI files
    if ai_queue:
        batch_result = process_ai_batch(ai_queue, db, config, low_conf_base,
                                        deletion_base, dry_run, safe_mode)
        for result in batch_result:
            if result:
                moved_deletion.append(result) if result[1].deletable else moved_low.append(result)
                if not dry_run and not safe_mode:
                    restore_list.append((result[0].path, result[2]))
    
    if progress:
        progress.stop()
    
    return moved_deletion, moved_low, restore_list

def process_ai_batch(file_infos: List[FileInfo], db: CleanWaveDB, config: dict,
                     low_conf_base: Path, deletion_base: Path, dry_run: bool, 
                     safe_mode: bool) -> List[Optional[Tuple[FileInfo, Decision, Optional[Path]]]]:
    """Process a batch of files through AI."""
    
    # Extract content for each file
    contents = []
    for f in file_infos:
        if f.size < MAX_AI_FILE_SIZE_BYTES:
            content = extract_content(f.path)
        else:
            content = "[File too large for content preview]"
        contents.append(content)
    
    # Try Groq first
    batch_result = call_groq_batch(file_infos, contents)
    
    results = []
    
    # Handle successful decisions
    for i, (file_info, decision) in enumerate(zip(file_infos, batch_result.decisions)):
        if decision:
            # Cache decision
            db.store_hash(file_info.path, file_info.size, file_info.mtime, 
                         f"ai_{decision.category}_{decision.confidence}")
            
            # Move based on decision
            if decision.deletable and decision.confidence >= config.get('confidence_threshold', 0.75):
                dest = move_to_review_folder(file_info, deletion_base, dry_run) if not safe_mode else None
                if not safe_mode and dest:
                    db.record_moved(file_info.path, dest, decision.category)
                results.append((file_info, decision, dest))
            elif decision.deletable and decision.confidence < config.get('confidence_threshold', 0.75):
                dest = move_to_review_folder(file_info, low_conf_base, dry_run) if not safe_mode else None
                if not safe_mode and dest:
                    db.record_moved(file_info.path, dest, f"low_conf_{decision.category}")
                results.append((file_info, decision, dest))
            else:
                # Not deletable - leave in place
                results.append(None)
    
    # Handle failed indices - try local Ollama fallback
    if batch_result.failed_indices and config.get('use_local_fallback', False):
        for i in batch_result.failed_indices:
            if i < len(file_infos):
                decision = call_local_ollama(file_infos[i], contents[i] if i < len(contents) else "")
                if decision:
                    if decision.deletable and decision.confidence >= config.get('confidence_threshold', 0.75):
                        dest = move_to_review_folder(file_infos[i], deletion_base, dry_run) if not safe_mode else None
                        if not safe_mode and dest:
                            db.record_moved(file_infos[i].path, dest, decision.category)
                        results.append((file_infos[i], decision, dest))
                    elif decision.deletable:
                        dest = move_to_review_folder(file_infos[i], low_conf_base, dry_run) if not safe_mode else None
                        if not safe_mode and dest:
                            db.record_moved(file_infos[i].path, dest, f"low_conf_{decision.category}")
                        results.append((file_infos[i], decision, dest))
                    else:
                        results.append(None)
                else:
                    results.append(None)
        else:
            for i in batch_result.failed_indices:
                results.append(None)
    
    return results

# ======================= CLEANUP UTILITIES =======================
def remove_empty_directories(start_path: Path, dry_run: bool) -> int:
    """Recursively remove empty directories."""
    removed = 0
    for root, dirs, files in os.walk(start_path, topdown=False):
        root_path = Path(root)
        try:
            if not any(root_path.iterdir()):
                if not dry_run:
                    root_path.rmdir()
                    console.print(f"[dim]Removed empty dir: {root_path}[/dim]")
                removed += 1
        except OSError:
            pass
    return removed

def find_large_old_files(files: List[FileInfo], threshold_gb: float, max_age_days: int) -> List[FileInfo]:
    """Find files exceeding size and age thresholds."""
    threshold_bytes = threshold_gb * 1024**3
    current_time = time.time()
    max_age_seconds = max_age_days * 86400
    
    large_files = []
    for f in files:
        age_seconds = current_time - f.mtime
        if f.size > threshold_bytes and age_seconds > max_age_seconds:
            large_files.append(f)
    
    return large_files

# ======================= REPORTING =======================
def generate_report(moved_deletion: List, moved_low: List, large_files: List, 
                    scan_dirs: List[str], duration_seconds: float, log_path: Path):
    """Generate comprehensive markdown report."""
    with open(log_path, 'w', encoding='utf-8') as f:
        f.write(f"# CleanWave Scan Report\n")
        f.write(f"**Date:** {datetime.datetime.now()}\n")
        f.write(f"**Duration:** {duration_seconds:.2f} seconds\n")
        f.write(f"**Scan Directories:** {', '.join(scan_dirs)}\n\n")
        
        f.write("## 📊 Summary\n")
        f.write(f"- **Moved to deletion_approval:** {len(moved_deletion)}\n")
        f.write(f"- **Moved to low_confidence_review:** {len(moved_low)}\n")
        f.write(f"- **Large old files found:** {len(large_files)}\n\n")
        
        f.write("## 🗑️ High Confidence Deletables\n")
        for file_info, decision, dest in moved_deletion[:100]:
            f.write(f"- **{file_info.path}** → `{dest}`\n")
            f.write(f"  - *Reason:* {decision.reason} (confidence: {decision.confidence:.0%})\n")
            f.write(f"  - *Category:* {decision.category}\n")
        if len(moved_deletion) > 100:
            f.write(f"\n*... and {len(moved_deletion) - 100} more files*\n")
        
        f.write("\n## ⚠️ Low Confidence (Review Manually)\n")
        for file_info, decision, dest in moved_low[:100]:
            f.write(f"- **{file_info.path}** → `{dest}`\n")
            f.write(f"  - *Reason:* {decision.reason} (confidence: {decision.confidence:.0%})\n")
        if len(moved_low) > 100:
            f.write(f"\n*... and {len(moved_low) - 100} more files*\n")
        
        if large_files:
            f.write("\n## 📦 Large Files (>1GB, untouched >1 year)\n")
            for file_info in large_files[:50]:
                size_gb = file_info.size / (1024**3)
                age_days = (time.time() - file_info.mtime) / 86400
                f.write(f"- **{file_info.path}** ({size_gb:.2f} GB, untouched {age_days:.0f} days)\n")
            if len(large_files) > 50:
                f.write(f"\n*... and {len(large_files) - 50} more files*\n")
    
    console.print(f"[green]✓ Report saved to {log_path}[/green]")

# ======================= COPY KEEPERS =======================
def copy_keepers_to_destination(scan_dirs: List[str], dest_base: Path, 
                                 exclude_paths: List[Path], dry_run: bool) -> int:
    """Copy all kept files (not in exclude_paths) to external destination."""
    copied = 0
    
    for src_dir in scan_dirs:
        src_path = Path(src_dir).expanduser()
        if not src_path.exists():
            continue
        
        for root, dirs, files in os.walk(src_path):
            root_path = Path(root)
            
            # Skip excluded directories
            dirs[:] = [d for d in dirs if not any(
                (root_path / d) in exclude_paths or 
                any((root_path / d).is_relative_to(excl) for excl in exclude_paths)
            )]
            
            for fname in files:
                fpath = root_path / fname
                
                # Skip excluded files
                if any(fpath in exclude_paths or fpath.is_relative_to(excl) for excl in exclude_paths):
                    continue
                
                # Compute relative path
                try:
                    rel = fpath.relative_to(src_path)
                except ValueError:
                    rel = Path(fpath.name)
                
                dest = dest_base / rel
                dest.parent.mkdir(parents=True, exist_ok=True)
                
                if not dry_run:
                    shutil.copy2(fpath, dest)
                copied += 1
                
                if copied % 1000 == 0:
                    console.print(f"[dim]Copied {copied} files...[/dim]")
    
    return copied

# ======================= MAIN =======================
def main():
    parser = argparse.ArgumentParser(description="CleanWave - Factory Reset Assistant")
    parser.add_argument("--scan-dirs", nargs="+", help="Directories to scan")
    parser.add_argument("--full-drive", action="store_true", help="Scan entire user home")
    parser.add_argument("--config", help="Path to config file (YAML or JSON)")
    parser.add_argument("--resume", action="store_true", help="Resume from last checkpoint")
    parser.add_argument("--dry-run", action="store_true", help="Simulate without moving files")
    parser.add_argument("--safe-mode", action="store_true", help="Don't move files, only report")
    parser.add_argument("--no-ai", action="store_true", help="Disable AI (rules only)")
    parser.add_argument("--copy-keepers-to", help="Copy kept files to external drive")
    parser.add_argument("--remove-empty-dirs", action="store_true", help="Remove empty directories after move")
    parser.add_argument("--find-large-files", action="store_true", help="Find large old files")
    parser.add_argument("--generate-restore-script", action="store_true", help="Generate restore script for moved files")
    
    args = parser.parse_args()
    
    # Load configuration
    config = {
        'ai_enabled': not args.no_ai,
        'confidence_threshold': DEFAULT_CONFIDENCE,
        'batch_size': 6,
        'parallel_hash_workers': 4,
        'use_local_fallback': False,
        'large_file_threshold_gb': 1,
        'large_file_age_days': 365,
        'remove_empty_dirs': args.remove_empty_dirs
    }
    
    if args.config and Path(args.config).exists():
        if args.config.endswith('.yaml') or args.config.endswith('.yml'):
            if YAML_AVAILABLE:
                with open(args.config) as f:
                    yaml_config = yaml.safe_load(f)
                    config.update(yaml_config)
        else:
            with open(args.config) as f:
                json_config = json.load(f)
                config.update(json_config)
    
    # Determine scan directories
    if args.full_drive:
        scan_dirs = [str(Path.home())]
    elif args.scan_dirs:
        scan_dirs = args.scan_dirs
    else:
        # Default: Downloads, Desktop, Documents
        home = str(Path.home())
        scan_dirs = [
            os.path.join(home, "Downloads"),
            os.path.join(home, "Desktop"),
            os.path.join(home, "Documents")
        ]
    
    # Setup paths
    cleanwave_dir = Path.home() / ".cleanwave"
    cleanwave_dir.mkdir(exist_ok=True)
    
    deletion_base = Path.home() / "deletion_approval"
    low_conf_base = Path.home() / "low_confidence_review"
    db_path = cleanwave_dir / "cleanwave.db"
    
    if not args.dry_run and not args.safe_mode:
        deletion_base.mkdir(exist_ok=True)
        low_conf_base.mkdir(exist_ok=True)
    
    # Initialize database
    db = CleanWaveDB(db_path)
    
    # Check external drive if copying
    if args.copy_keepers_to:
        dest_path = Path(args.copy_keepers_to).expanduser()
        mounted, msg = check_mount(dest_path)
        if not mounted:
            console.print(f"[red]ERROR: Cannot write to {dest_path}: {msg}[/red]")
            return
        has_space, free_gb = check_disk_space(dest_path, 5.0)
        if not has_space:
            console.print(f"[yellow]Warning: Only {free_gb:.1f}GB free on {dest_path}[/yellow]")
            if not Confirm.ask("Continue anyway?", default=False):
                return
    
    # Check local disk space
    has_space, free_gb = check_disk_space(Path.home(), 5.0)
    if not has_space:
        console.print(f"[yellow]Warning: Only {free_gb:.1f}GB free in home directory[/yellow]")
        if not Confirm.ask("Continue anyway?", default=False):
            return
    
    # Display banner
    if RICH_AVAILABLE:
        console.print(Panel.fit("[bold cyan]🧹 CleanWave - Factory Reset Assistant[/bold cyan]", border_style="cyan"))
        console.print(f"\n[cyan]Scanning:[/cyan] {', '.join(scan_dirs)}")
        console.print(f"[cyan]AI:[/cyan] {'Enabled' if config['ai_enabled'] else 'Disabled'}")
        console.print(f"[cyan]Mode:[/cyan] {'Safe (no moves)' if args.safe_mode else 'Normal'}")
        if args.dry_run:
            console.print("[yellow]DRY RUN - No files will be moved[/yellow]")
    else:
        print(f"CleanWave - Scanning: {', '.join(scan_dirs)}")
    
    start_time = time.time()
    scan_id = db.begin_scan(config)
    
    # Collect files
    console.print("\n[cyan]📁 Collecting files...[/cyan]")
    exclusions = get_os_exclusions()
    all_files = collect_files(scan_dirs, exclusions, db, not args.dry_run, None)
    console.print(f"[green]✓ Found {len(all_files)} files[/green]")
    
    # Find large files if requested
    if args.find_large_files:
        large_files = find_large_old_files(
            all_files,
            config.get('large_file_threshold_gb', 1),
            config.get('large_file_age_days', 365)
        )
        if large_files:
            console.print(f"\n[yellow]📦 Found {len(large_files)} large files (>1GB, untouched >1 year)[/yellow]")
            for f in large_files[:20]:
                size_gb = f.size / (1024**3)
                age_days = (time.time() - f.mtime) / 86400
                console.print(f"  • {f.path.name} ({size_gb:.2f} GB, {age_days:.0f} days old)")
            if len(large_files) > 20:
                console.print(f"  ... and {len(large_files) - 20} more")
    
    # Process files
    whitelist = Whitelist(
        DEFAULT_WHITELIST_PATTERNS + config.get('extra_whitelist_patterns', []),
        config.get('whitelist_paths', [])
    )
    
    moved_deletion, moved_low, restore_list = process_files_advanced(
        all_files, deletion_base, low_conf_base, db, whitelist, config,
        args.dry_run, args.safe_mode, args.resume
    )
    
    # Generate restore script if requested
    if args.generate_restore_script and restore_list and not args.dry_run:
        generate_restore_script(restore_list, cleanwave_dir)
    
    # Remove empty directories
    if config.get('remove_empty_dirs', False) and not args.dry_run:
        console.print("\n[cyan]🗑️ Removing empty directories...[/cyan]")
        removed = remove_empty_directories(Path.home(), args.dry_run)
        console.print(f"[green]✓ Removed {removed} empty directories[/green]")
    
    # Copy keepers to external drive
    if args.copy_keepers_to:
        console.print(f"\n[cyan]💾 Copying kept files to {args.copy_keepers_to}...[/cyan]")
        exclude_paths = [deletion_base, low_conf_base]
        copied = copy_keepers_to_destination(scan_dirs, Path(args.copy_keepers_to), exclude_paths, args.dry_run)
        console.print(f"[green]✓ Copied {copied} files[/green]")
    
    # Generate report
    duration = time.time() - start_time
    log_path = cleanwave_dir / f"cleanwave_report_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.md"
    generate_report(moved_deletion, moved_low, large_files if args.find_large_files else [], 
                   scan_dirs, duration, log_path)
    
    # Final summary
    db.end_scan(scan_id, len(all_files), len(moved_deletion) + len(moved_low))
    
    console.print("\n[bold green]✅ CleanWave Complete![/bold green]")
    console.print(f"   ⏱️  Duration: {duration:.1f} seconds")
    console.print(f"   🗑️  Moved to deletion_approval: {len(moved_deletion)}")
    console.print(f"   ⚠️  Moved to low_confidence_review: {len(moved_low)}")
    console.print(f"   📁 Review folders: {deletion_base}\n                 {low_conf_base}")
    console.print(f"   📄 Report: {log_path}")
    
    # Notification
    try:
        notification.notify(
            title="CleanWave Complete",
            message=f"Moved {len(moved_deletion) + len(moved_low)} files for review",
            app_name="CleanWave",
            timeout=5
        )
    except:
        pass
    
    db.close()

if __name__ == "__main__":
    main()